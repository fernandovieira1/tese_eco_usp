from pathlib import Path
import json,hashlib,math,collections,datetime,zipfile
import openpyxl
HERE=Path(__file__).resolve().parent;APP=HERE.parent
ROOT=next(p for p in APP.parents if (p/'2_INTI/RESEARCH_PROTOCOL.md').exists())
source=json.loads((HERE/'candidatas_origem.json').read_text(encoding='utf8'))
meta=json.loads((HERE/'metadata_fe.json').read_text(encoding='utf8'))
rows_by_sheet=json.loads((HERE/'excel_fe_payload.json').read_text(encoding='utf8'))
hashfile=lambda p:hashlib.sha256(Path(p).read_bytes()).hexdigest()
assert hashfile(APP/'RESULTADOS_DASH_H1.xlsx')==source['sha256'],'Excel anterior foi alterado.'
assert hashfile(ROOT/meta['data']['csv'])==meta['data']['csv_sha256']
for file,digest in meta['code_hashes'].items():assert hashfile(ROOT/file)==digest
path=APP/'RESULTADOS_DASH_H1_FE.xlsx'
wb=openpyxl.load_workbook(path,data_only=False)
assert wb.sheetnames==list(rows_by_sheet)
count=0
for name,expected in rows_by_sheet.items():
    sh=wb[name]
    assert sh.max_row==len(expected) and sh.max_column==len(expected[0]),name
    assert sh.freeze_panes=='A2' and len(sh.tables)==1
    assert next(iter(sh.tables.values())).autoFilter is not None
    for i,row in enumerate(expected,1):
        for j,value in enumerate(row,1):
            cell=sh.cell(i,j);actual=cell.value
            assert cell.data_type!='e',(name,i,j)
            if value=='':value=None
            if isinstance(value,(int,float)) and not isinstance(value,bool):
                assert isinstance(actual,(int,float)) and math.isfinite(actual),(name,i,j)
                assert math.isclose(actual,value,rel_tol=1e-13,abs_tol=1e-15),(name,i,j,actual,value)
            else:assert actual==value,(name,i,j,str(actual)[:80],str(value)[:80])
            count+=1
results=list(wb['Resultados'].values)[1:]
assert len(results)==12 and len({row[0] for row in results})==12
assert set(collections.Counter(row[1] for row in results).values())=={3}
details=list(wb['Detalhes'].values)[1:]
assert len(details)==36 and len({row[0] for row in details})==36
for mid in {row[0] for row in results}:
    group=[row for row in details if row[1]==mid]
    assert len(group)==3 and len({row[4] for row in group})==len({row[15] for row in group})==1
    for row in group:assert row[13]==row[4]-row[12] and row[14]==row[5]-1
absorption=list(wb['Absorção'].values)[1:]
absorbed=[row for row in absorption if row[4].startswith('Absorvido')]
assert len(absorbed)==36
coef=list(wb['Coeficientes'].values)[1:]
unidentified=[row for row in coef if row[4]=='Regressor absorvido']
assert len(unidentified)==36
assert all(all(value is None for value in row[5:11]) for row in unidentified)
for row in details:
    estimated=[c for c in coef if c[0]==row[0] and c[4]!='Regressor absorvido']
    assert len(estimated)==row[12]
assert len({(row[0],row[3]) for row in coef})==len(coef)
assert all(r['economic_coefficient_inference_finite'] for r in json.loads((HERE/'diagnosticos_fe.json').read_text(encoding='utf8')))
with zipfile.ZipFile(path) as z:assert z.testzip() is None
verification={'status':'VERIFICADO','new_workbook':str(path),'sha256':hashfile(path),'source_sha256_before_after':source['sha256'],
 'preserved_previous_excel':True,'preserved_source_csv_and_code':True,'twfe_models':12,'pooled_comparisons':24,'total_fits':36,
 'absorbed_terms':36,'duplicate_coefficients':0,'compared_cells':count,'sheets':wb.sheetnames,
 'pairwise_samples_identical':True,'date':datetime.datetime.now(datetime.timezone(datetime.timedelta(hours=-3))).isoformat()}
(HERE/'verificacao_fe.json').write_text(json.dumps(verification,ensure_ascii=False,indent=2),encoding='utf8')
audit=ROOT/'2_INTI/audit/fe_candidatas_20261018';audit.mkdir(parents=True,exist_ok=True)
(audit/'verificacao_fe.json').write_text(json.dumps(verification,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps(verification,ensure_ascii=False))
