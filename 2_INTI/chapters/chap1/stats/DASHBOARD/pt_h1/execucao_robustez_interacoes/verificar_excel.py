from pathlib import Path
import openpyxl,json,math,hashlib,zipfile,datetime
HERE=Path(__file__).resolve().parent;APP=HERE.parent
ROOT=next(p for p in APP.parents if (p/'2_INTI/RESEARCH_PROTOCOL.md').exists())
plan=json.loads((HERE/'plano_robustez.json').read_text(encoding='utf8'))
records=json.loads((HERE/'resultados_robustez.json').read_text(encoding='utf8'))
expected=json.loads((HERE/'excel_payload.json').read_text(encoding='utf8'))
hashfile=lambda p:hashlib.sha256(Path(p).read_bytes()).hexdigest()
for name,digest in plan['source_xlsx_sha256'].items():assert hashfile(APP/name)==digest,'Excel anterior modificado: '+name
for name,digest in plan['code_sha256'].items():assert hashfile(ROOT/name)==digest
metadata=json.loads((HERE/'metadata_robustez.json').read_text(encoding='utf8'))
assert hashfile(ROOT/metadata['data']['csv'])==plan['csv_sha256']
path=APP/'RESULTADOS_DASH_H1_FE_ROBUSTEZ.xlsx';wb=openpyxl.load_workbook(path,data_only=False)
assert wb.sheetnames==list(expected)
cells=0
for name,rows in expected.items():
    sh=wb[name];assert sh.max_row==len(rows) and sh.max_column==len(rows[0])
    assert sh.freeze_panes=='A2' and len(sh.tables)==1
    assert next(iter(sh.tables.values())).autoFilter is not None
    for i,row in enumerate(rows,1):
        for j,value in enumerate(row,1):
            actual=sh.cell(i,j).value;assert sh.cell(i,j).data_type!='e'
            if value=='':value=None
            if isinstance(value,(int,float)) and not isinstance(value,bool):
                assert isinstance(actual,(int,float)) and math.isfinite(actual),(name,i,j)
                assert math.isclose(actual,value,rel_tol=1e-13,abs_tol=1e-15),(name,i,j,actual,value)
            else:assert actual==value,(name,i,j,str(actual)[:80],str(value)[:80])
            cells+=1
rows=list(wb['Robustez'].values)[1:];assert len(rows)==4
assert {r[0] for r in rows}=={'M042','M092','M142','M192'}
assert all(r[2]==2655 and r[11]==137 and r[12]==23 and r[13]==22 for r in rows)
assert all(r[5]>0.05 and r[9]<0<r[10] for r in rows)
models=list(wb['Modelos'].values)[1:];assert len(models)==16 and len({r[0] for r in models})==16
assert len({r[14] for r in models})==1
for mid in plan['ids']:
    pair=[r for r in models if r[1]==mid];assert len(pair)==2
    assert pair[0][9:12]==pair[1][9:12]
coeff=list(wb['Coeficientes'].values)[1:]
assert len({(r[0],r[4]) for r in coeff})==len(coeff)
negative_variance=[r for r in coeff if r[13] is not None and r[13]<0]
assert negative_variance
assert all(r[6]=='Dummy FE' and all(x is None for x in r[8:13]) for r in negative_variance)
for r in records:
    assert r['sample_id']==plan['common_sample_sha256']
    assert not r['covariance_negative_diagonal_economic']
    assert r['covariance_economic_psd_tol1e8']
    if r['candidate'] in plan['focus'] and r['fit_spec']['cov']=='cluster_twoway':
        inter=next(c for c in r['coef'] if ' × ' in c['term'])
        assert inter['variance']>0 and inter['se'] is not None
with zipfile.ZipFile(path) as z:assert z.testzip() is None
verification={'status':'VERIFICADO','new_excel':str(path),'sha256':hashfile(path),'previous_workbooks_preserved':True,
 'csv_code_preserved':True,'interactions':4,'twfe_specifications':8,'fits':16,'common_N':2655,'entity_clusters':137,'year_clusters':23,
 'two_way_t_df':22,'nonpositive_interaction_variances':0,'all_interactions_CI95_include_zero':True,
 'negative_nuisance_variances_recorded':len(negative_variance),'compared_cells':cells,'duplicates':0,'sheets':wb.sheetnames,
 'date':datetime.datetime.now(datetime.timezone(datetime.timedelta(hours=-3))).isoformat()}
(HERE/'verificacao.json').write_text(json.dumps(verification,ensure_ascii=False,indent=2),encoding='utf8')
audit=ROOT/'2_INTI/audit/robustez_interacoes_20261018';audit.mkdir(parents=True,exist_ok=True)
(audit/'verificacao.json').write_text(json.dumps(verification,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps(verification,ensure_ascii=False))
